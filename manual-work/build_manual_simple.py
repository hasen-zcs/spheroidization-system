# -*- coding: utf-8 -*-
"""生成《球化监测系统用户培训手册（精简版）》docx——供一线工人操作使用。

主线：每天五步（看工位 → 录数据 → 查结果 → 处理异常 → 看统计），
语言尽量简单，配关键截图与快速查表。
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = r"D:\项目组学习\项目工作\东贝生产线\系统架构\1.2\球化监测系统 (1)\spheroidization-system"
SHOT = os.path.join(ROOT, "manual-work")
OUT = os.path.join(ROOT, "docx", "球化监测系统用户培训手册（精简版）.docx")

FONT = "微软雅黑"
BLUE = RGBColor(0x00, 0x33, 0x99)
BLUE2 = RGBColor(0x00, 0x50, 0xB3)
GREY = RGBColor(0x77, 0x77, 0x77)
DARK = RGBColor(0x1A, 0x2A, 0x3A)
RED = RGBColor(0xC6, 0x28, 0x28)
ORANGE = RGBColor(0xD9, 0x77, 0x06)
GREEN = RGBColor(0x16, 0x80, 0x3C)

BS = 11  # 正文加大一号，方便车间阅读


def set_cjk(rpr, name=FONT):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(attr), name)


def style_run(run, size=BS, bold=False, color=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    set_cjk(run._element.get_or_add_rPr(), name)


def add_para(doc, text="", size=BS, bold=False, color=None, align=None,
             space_before=0, space_after=6, indent=None, line=1.4):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    if text:
        style_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_rich(doc, parts, align=None, space_after=6, indent=None, line=1.4):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    for text, kw in parts:
        style_run(p.add_run(text), **kw)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    sizes = {1: 17, 2: 13.5, 3: 12}
    colors = {1: BLUE, 2: BLUE2, 3: DARK}
    run = p.add_run(text)
    style_run(run, size=sizes.get(level, 12), bold=True, color=colors.get(level, DARK))
    pf = p.paragraph_format
    pf.space_before = Pt(18 if level == 1 else 10)
    pf.space_after = Pt(8 if level == 1 else 6)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.4
    pf.left_indent = Cm(0.5)
    style_run(p.add_run("• "), size=BS, color=BLUE2)
    if bold_prefix:
        style_run(p.add_run(bold_prefix), size=BS, bold=True)
    style_run(p.add_run(text), size=BS)
    return p


def add_step(doc, num, text, note=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(5)
    pf.line_spacing = 1.4
    pf.left_indent = Cm(0.3)
    style_run(p.add_run("第 %d 步　" % num), size=BS, bold=True, color=BLUE)
    style_run(p.add_run(text), size=BS)
    if note:
        style_run(p.add_run("（" + note + "）"), size=10, color=GREY)
    return p


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, size=10.5, color=None, align=None, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_after = Pt(2)
    pf.space_before = Pt(2)
    pf.line_spacing = 1.3
    if align is not None:
        p.alignment = align
    style_run(p.add_run(text), size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)


def add_table(doc, headers, rows, widths=None, header_fill="003399",
              zebra=True, font_size=10.5, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, bold=True, size=font_size,
                  color=RGBColor(0xFF, 0xFF, 0xFF),
                  align=WD_ALIGN_PARAGRAPH.CENTER, fill=header_fill)
    for r, row in enumerate(rows):
        cells = table.add_row().cells
        fill = "F2F7FC" if (zebra and r % 2 == 1) else None
        for i, val in enumerate(row):
            cell_text(cells[i], str(val), size=font_size,
                      bold=(first_col_bold and i == 0), fill=fill)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    return table


def add_figure(doc, shot_file, caption, width_cm=15.4):
    path = os.path.join(SHOT, shot_file)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    style_run(cap.add_run(caption), size=9.5, color=GREY)


def add_tip(doc, title, text, fill="EFF6FF", bar_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    cell.width = Cm(15.4)
    shade_cell(cell, fill)
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.35
    style_run(p1.add_run(title + "："), size=10.5, bold=True, color=bar_color)
    style_run(p1.add_run(text), size=10.5)
    add_para(doc, "", size=4, space_after=2)


def add_page_number(paragraph):
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = "PAGE"
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)
    style_run(run, size=9, color=GREY)


# ============================== 文档 ==============================

doc = Document()

# 页面：A4
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)

# Normal 样式
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(BS)
set_cjk(normal._element.get_or_add_rPr())

# 页脚
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_run(fp.add_run("球化监测系统一线操作简明手册（精简版）　·　第 "), size=9, color=GREY)
add_page_number(fp)
style_run(fp.add_run(" 页"), size=9, color=GREY)

# ---------------- 封面 ----------------
for _ in range(5):
    add_para(doc, "", space_after=12)
add_para(doc, "球化监测系统", size=30, bold=True, color=BLUE,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
add_para(doc, "一线操作简明手册", size=22, bold=True, color=DARK,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
add_para(doc, "（精简版 · 供一线操作工日常使用）", size=13, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=60)

info = doc.add_table(rows=4, cols=2)
info.style = "Table Grid"
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info_rows = [("文档版本", "V1.0（精简版 · 一线操作）"),
             ("适用对象", "浇注工位操作工、班组长"),
             ("配套文档", "《球化监测系统用户培训手册（详细版）》——完整功能说明，供项目存档"),
             ("编制单位", "东贝生产线 · 系统架构项目组（2026 年 10 月）")]
for i, (k, v) in enumerate(info_rows):
    cell_text(info.rows[i].cells[0], k, bold=True, size=10.5,
              align=WD_ALIGN_PARAGRAPH.CENTER, fill="E6F2FF")
    cell_text(info.rows[i].cells[1], v, size=10.5)
    info.rows[i].cells[0].width = Cm(3.6)
    info.rows[i].cells[1].width = Cm(10.0)
add_para(doc, "", space_after=16)
add_para(doc, "东贝生产线 · 系统架构项目组", size=11, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ---------------- 使用说明 ----------------
add_heading(doc, "使用说明", 1)
add_rich(doc, [
    ("这本手册只讲你每天要做的操作。", {"size": BS, "bold": True}),
    ("一句话记住分工：", {"size": BS}),
    ("机器负责“检”，你负责“看结果、录数据、处理异常”。", {"size": BS, "bold": True, "color": BLUE}),
])

add_heading(doc, "系统共 5 个页面，各管什么事", 2)
add_table(doc,
          ["页面（入口）", "你在上面做什么"],
          [
              ["浇注工位监控（系统首页）", "看摄像头画面、最新 5 条检测数据"],
              ["首页（统计）", "看总记录数、异常数、已处理/未处理"],
              ["球化记录", "新增检测记录、查询所有记录"],
              ["异常记录", "查看异常记录、填写复核、标记处理"],
              ["检测标准", "查看当前标准值（一般由工艺人员调整）"],
          ],
          widths=[6.0, 9.4], first_col_bold=True)
add_para(doc, "", space_after=2)

add_heading(doc, "怎么打开系统", 2)
add_step(doc, 1, "打开浏览器（推荐 Edge 或 Chrome 最新版）。")
add_step(doc, 2, "在地址栏输入系统地址，回车。")
add_step(doc, 3, "不需要登录，页面打开就是“浇注工位监控”页。")
add_tip(doc, "提示", "页面右上角的“系统正常”按钮，表示系统连接正常。"
        "如果显示“系统离线”，请联系 IT 人员。")

# ---------------- 1 每天五步 ----------------
add_heading(doc, "1　每天五步操作", 1)

# ---- 第 1 步 ----
add_heading(doc, "第 1 步：看工位（系统首页）", 2)
add_para(doc, "接班后的第一件事。打开系统，看三样东西：")
add_bullet(doc, "两个摄像头画面是否正常（显示“等待视频流接入”表示摄像头还没连上，不用管）；")
add_bullet(doc, "最近 5 条检测数据的“球化结果”一列：正常还是异常；")
add_bullet(doc, "最下面蓝色状态栏：最新一条记录是正常/异常、已处理还是未处理。")
add_para(doc, "看到“异常”而且还没处理 → 直接跳到下面的“第 4 步”。")
add_figure(doc, "01_monitor.png", "图 1　浇注工位监控首页（摄像头画面、最新检测数据、底部状态栏）")

# ---- 第 2 步 ----
add_heading(doc, "第 2 步：录入一条新的检测记录", 2)
add_para(doc, "入口：在首页点“＋ 新增一行数据”；或者打开“球化记录”页面点“新增记录”。")
add_figure(doc, "04_records_create.png", "图 2　新增球化记录表单（按表格内容填写即可）")
add_para(doc, "填写对照表（照着填就行）：")
add_table(doc,
          ["字段", "填什么", "说明"],
          [
              ["检测时间", "本次检测的日期时间", "必填"],
              ["球化开始时间 / 结束时间", "球化反应开始、结束时刻", "必填"],
              ["实际球化时间", "实际球化时长（秒）", "必填"],
              ["球化温度 / 测温时间", "铁水温度、测温时刻", ""],
              ["进线开始 / 结束时间", "进线开始、结束时刻", ""],
              ["进线总时长", "进线总时长（秒）", ""],
              ["球化重量", "铁水重量", ""],
              ["机器判定异常", "机器判为异常时，勾选它", ""],
              ["球化图片", "上传检测图片", "只能选 JPG"],
          ],
          widths=[4.6, 7.4, 3.4])
add_para(doc, "", space_after=2)
add_para(doc, "填完点“提交记录”，新记录马上出现在列表最上面。")
add_tip(doc, "注意", "图片只能选 JPG 格式，选别的格式会提示“请选择有效的 JPG 图片”。"
        "机器判定为异常时，一定要勾选“机器判定异常”，这样它才会出现在异常记录页。",
      fill="FFF4E5", bar_color=ORANGE)

# ---- 第 3 步 ----
add_heading(doc, "第 3 步：查看检测结果", 2)
add_para(doc, "“球化记录”页面是全部记录的列表（每页 20 条，可翻页）。")
add_figure(doc, "03_records.png", "图 3　球化记录列表（查询条件、记录表格、分页）")
add_para(doc, "想查某些记录时：")
add_step(doc, 1, "“检测结果”下拉框选：全部 / 正常 / 异常。")
add_step(doc, 2, "“处理状态”下拉框选：全部 / 已处理 / 未处理。")
add_step(doc, 3, "点“查询”，列表就只剩符合条件的记录。")
add_para(doc, "看列表时重点看“标准/实际球化时长”这一列：前面是标准值，后面是实际值，"
              "一眼就能看出这次球化达不达标。想细看某一条，点该行“查看详情”。")

# ---- 第 4 步 ----
add_heading(doc, "第 4 步：处理异常记录（重点）", 2)
add_para(doc, "入口：左侧菜单“异常记录”，或顶部导航栏的“异常记录”。"
              "这个页面自动只列机器判“异常”的记录。")
add_figure(doc, "05_abnormal.png", "图 4　异常记录页面（自动只列出异常记录）")
add_para(doc, "找到“未处理”的异常记录，点“查看详情”，在详情页做三件事：")
add_step(doc, 1, "看曲线：详情页最上面的“球化亮光分析曲线”，帮助判断异常原因。")
add_step(doc, 2, "人工复核：选“复核结果”，填“复核人”（必填）、“复核说明”（必填），点“提交复核”。")
add_step(doc, 3, "标记处理：现场处置完成后，填“处理人”，点“标记为已处理”。")
add_figure(doc, "08_detail_unprocessed.png", "图 5　未处理的记录详情（底部可填处理人、点“标记为已处理”）")
add_figure(doc, "07_detail_processed.png", "图 6　处理完成后的详情（复核记录、处理人、处理时间都留痕）")

add_para(doc, "", space_after=2)
add_rich(doc, [
    ("三条红线（务必记住）：", {"size": BS, "bold": True, "color": RED}),
], space_after=4)
add_bullet(doc, "机器检测结果任何人不能改；你的复核结论单独保存，两者并存、不互相覆盖。")
add_bullet(doc, "“复核”和“标记为已处理”是两件独立的事，各做各的，别混。")
add_bullet(doc, "点了“标记为已处理”就不能撤销——点之前一定确认异常已经实际处理好。")

# ---- 第 5 步 ----
add_heading(doc, "第 5 步：看统计（首页）", 2)
add_para(doc, "打开“首页”，四个数字一眼看全局：")
add_bullet(doc, "总记录数：一共检了多少次；", bold_prefix=None)
add_bullet(doc, "异常记录：机器判异常的次数（红色）；")
add_bullet(doc, "已处理：处理完的次数（绿色）；")
add_bullet(doc, "未处理：还没处理的次数（橙色）。")
add_para(doc, "班组长重点盯一个数：让“未处理”尽快变成 0。")
add_figure(doc, "02_home.png", "图 7　首页统计概览（四个统计数字 + 最近检测记录）")

# ---------------- 2 快速查表 ----------------
add_heading(doc, "2　遇到问题快速查表", 1)
add_table(doc,
          ["遇到什么情况", "怎么办"],
          [
              ["摄像头窗口显示“等待视频流接入”", "摄像头还没连上，不用操作，接上后会自动显示画面"],
              ["右上角显示“系统离线”或一直“连接中”", "后端服务没启动，联系 IT 人员"],
              ["页面提示“无法连接后端服务”", "联系 IT 人员检查服务"],
              ["新增时提示“请选择有效的 JPG 图片”", "图片格式不对，换成 JPG 图片再传"],
              ["提交复核时提示“请输入复核人”“请输入复核说明”", "这两项是必填的，填上再点提交"],
              ["机器判异常，能不能改成正常？", "不能。机器结果永久保留，复核结论单独保存"],
              ["误点了“标记为已处理”能撤销吗？", "不能，系统不可回退，请联系工艺技术人员/IT 处理"],
              ["想改球化时间标准", "“检测标准”页面，一般由工艺技术人员操作"],
          ],
          widths=[7.2, 8.2])

# ---------------- 3 找谁 ----------------
add_heading(doc, "3　出问题了找谁", 1)
add_table(doc,
          ["情况", "找谁"],
          [
              ["页面打不开 / 数据不刷新 / 摄像头没画面", "IT 系统管理员"],
              ["对检测结果有争议 / 要调整球化标准", "工艺 / 质量技术人员"],
              ["其他系统问题", "先报班组长，再由班组长联系项目系统管理员"],
          ],
          widths=[9.4, 6.0], first_col_bold=True)
add_para(doc, "（具体姓名和联系电话可由班组长填写后发到班组。）", size=10, color=GREY)

add_para(doc, "", space_after=6)
add_para(doc, "需要查更详细的说明，请看配套文档《球化监测系统用户培训手册（详细版）》。",
         size=10.5, color=GREY)
add_para(doc, "—— 手册完 ——", size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print("saved:", OUT, os.path.getsize(OUT), "bytes")
